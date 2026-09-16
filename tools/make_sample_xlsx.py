# -*- coding: utf-8 -*-
"""
데이터 분석 페이지를 시험해 볼 예시 엑셀과 입력 양식을 만듭니다.
  python tools/make_sample_xlsx.py
→ samples/은혜푸드마켓_2026_기부물품_예시.xlsx
→ samples/기부물품_입력양식.xlsx
"""
import os, random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

random.seed(2026)
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), "samples")
os.makedirs(OUT, exist_ok=True)

HEAD_FILL = PatternFill("solid", fgColor="E6F4EC")
HEAD_FONT = Font(bold=True, color="145A36")

def style_header(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill = HEAD_FILL; c.font = HEAD_FONT; c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"

donors = [
    ("㈜미트리", "기업·제조사"), ("동탄농협 하나로마트", "유통·마트"), ("이마트 동탄점", "유통·마트"),
    ("경기나눔푸드뱅크", "공공 배분"), ("전국푸드뱅크", "공공 배분"), ("동탄명성교회", "종교단체"),
    ("동탄반석교회", "종교단체"), ("운평장로교회", "종교단체"), ("오뚜기 화성공장", "기업·제조사"),
    ("CJ제일제당", "기업·제조사"), ("개인 기탁", "개인"), ("동탄시티병원", "기업·제조사"),
    ("화성시 공공급식지원센터", "공공 배분"), ("하루와규 동탄점", "기업·제조사"),
]
items = [
    ("쌀 10kg", "쌀·곡류", 32000), ("현미 4kg", "쌀·곡류", 18000), ("참치캔 6입", "가공식품·통조림", 9800),
    ("햄 통조림", "가공식품·통조림", 4500), ("라면 5입", "라면·즉석식품", 4200), ("즉석밥 12입", "라면·즉석식품", 13500),
    ("우유 1L", "음료·유제품", 2800), ("두유 24입", "음료·유제품", 21000), ("사과 5kg", "신선식품", 25000),
    ("양파 3kg", "신선식품", 6500), ("돼지고기 1kg", "신선식품", 14000), ("세제 3L", "생활·위생용품", 9900),
    ("화장지 30롤", "생활·위생용품", 17000), ("생리대 세트", "생활·위생용품", 8500), ("김 20봉", "가공식품·통조림", 12000),
    ("식용유 1.8L", "가공식품·통조림", 7900), ("마스크 50매", "생활·위생용품", 6000), ("과자 세트", "기타", 5500),
]
dongs = [f"동탄{i}동" for i in range(1, 10)] + ["화산동"]

# ---------- 1) 기부물품 접수·배분 ----------
wb = Workbook()
ws = wb.active; ws.title = "기부물품접수배분"
ws.append(["날짜", "구분", "기탁처", "기탁처유형", "품목", "분류", "수량", "단가(원)", "금액(원)", "비고"])
start = dt.date(2026, 1, 2)
rows = []
for day in range(0, 240):
    d = start + dt.timedelta(days=day)
    if d.weekday() >= 5:
        continue
    for _ in range(random.choice([0, 1, 1, 2])):
        donor, dtype = random.choice(donors)
        name, cat, price = random.choice(items)
        qty = random.choice([10, 20, 24, 30, 48, 50, 60, 100, 120])
        rows.append([d, "기탁", donor, dtype, name, cat, qty, price, qty * price, ""])
    for _ in range(random.choice([1, 2, 2, 3])):
        name, cat, price = random.choice(items)
        qty = random.choice([5, 8, 10, 12, 15, 20, 24, 30])
        rows.append([d, "배분", random.choice(dongs) + " 이용가구", "이용자", name, cat, qty, price, qty * price, "푸드마켓" if random.random() < .6 else "그냥드림"])
# 미트리 대량 기탁 (실제 보도)
rows.append([dt.date(2026, 5, 14), "기탁", "㈜미트리", "기업·제조사", "자사 제품(닭가슴살 등)", "가공식품·통조림", 1240, 3500, 1240 * 3500, "가정의 달 기탁 (하이뉴스 보도)"])
rows.sort(key=lambda r: r[0])
for r in rows:
    ws.append(r)
for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
    row[0].number_format = "yyyy-mm-dd"
for row in ws.iter_rows(min_row=2, min_col=7, max_col=9):
    for c in row:
        c.number_format = "#,##0"
style_header(ws, [12, 8, 22, 14, 22, 18, 8, 10, 12, 26])

# ---------- 2) 이용 현황 ----------
ws2 = wb.create_sheet("이용현황")
ws2.append(["월", "동", "이용가구", "이용인원", "이용건수", "그냥드림이용"])
for m in range(1, 9):
    for dong in dongs:
        hh = random.randint(18, 60)
        ws2.append([f"2026-{m:02d}", dong, hh, int(hh * random.uniform(1.4, 2.1)), int(hh * random.uniform(1.5, 2.4)), random.randint(10, 90)])
style_header(ws2, [10, 10, 10, 10, 10, 12])

# ---------- 3) 재정 수입·지출 ----------
ws3 = wb.create_sheet("재정수입지출")
ws3.append(["날짜", "수입/지출", "계정과목", "적요", "금액(원)"])
acc_in = [("보조금", "화성시 기부식품 제공사업 보조금"), ("후원금", "정기 후원"), ("후원금", "일시 후원"), ("사업수입", "협동조합 사업수입")]
acc_out = [("인건비", "직원 급여"), ("운영비", "사무실 임차료"), ("운영비", "차량 유류·정비"), ("운영비", "전기·수도·통신"), ("사업비", "부족 물품 구입"), ("사업비", "행사·홍보")]
for m in range(1, 9):
    d = dt.date(2026, m, 5)
    ws3.append([d, "수입", "보조금", acc_in[0][1], 8000000])
    ws3.append([d + dt.timedelta(days=3), "수입", "후원금", acc_in[1][1], random.randint(1500000, 3200000)])
    if random.random() < .5:
        ws3.append([d + dt.timedelta(days=10), "수입", "후원금", acc_in[2][1], random.randint(300000, 1500000)])
    ws3.append([d + dt.timedelta(days=12), "수입", "사업수입", acc_in[3][1], random.randint(600000, 1400000)])
    ws3.append([dt.date(2026, m, 25), "지출", "인건비", acc_out[0][1], 6000000])
    ws3.append([dt.date(2026, m, 25), "지출", "운영비", acc_out[1][1], 1500000])
    ws3.append([dt.date(2026, m, 20), "지출", "운영비", acc_out[2][1], random.randint(300000, 700000)])
    ws3.append([dt.date(2026, m, 18), "지출", "운영비", acc_out[3][1], random.randint(250000, 500000)])
    ws3.append([dt.date(2026, m, 15), "지출", "사업비", acc_out[4][1], random.randint(800000, 2600000)])
    if random.random() < .4:
        ws3.append([dt.date(2026, m, 16), "지출", "사업비", acc_out[5][1], random.randint(200000, 900000)])
for row in ws3.iter_rows(min_row=2, min_col=1, max_col=1):
    row[0].number_format = "yyyy-mm-dd"
for row in ws3.iter_rows(min_row=2, min_col=5, max_col=5):
    row[0].number_format = "#,##0"
style_header(ws3, [12, 10, 12, 30, 14])

p1 = os.path.join(OUT, "은혜푸드마켓_2026_기부물품_예시.xlsx")
wb.save(p1)

# ---------- 입력 양식 ----------
tw = Workbook()
t = tw.active; t.title = "기부물품접수배분"
t.append(["날짜", "구분", "기탁처", "기탁처유형", "품목", "분류", "수량", "단가(원)", "금액(원)", "비고"])
t.append([dt.date(2026, 1, 5), "기탁", "동탄농협 하나로마트", "유통·마트", "쌀 10kg", "쌀·곡류", 20, 32000, 640000, ""])
t.append([dt.date(2026, 1, 6), "배분", "동탄3동 이용가구", "이용자", "쌀 10kg", "쌀·곡류", 5, 32000, 160000, "푸드마켓"])
t["A2"].number_format = t["A3"].number_format = "yyyy-mm-dd"
style_header(t, [12, 8, 22, 14, 22, 18, 8, 10, 12, 26])
t2 = tw.create_sheet("재정수입지출")
t2.append(["날짜", "수입/지출", "계정과목", "적요", "금액(원)"])
t2.append([dt.date(2026, 1, 5), "수입", "보조금", "화성시 보조금", 8000000])
t2.append([dt.date(2026, 1, 25), "지출", "인건비", "직원 급여", 6000000])
t2["A2"].number_format = t2["A3"].number_format = "yyyy-mm-dd"
style_header(t2, [12, 10, 12, 30, 14])
g = tw.create_sheet("작성안내")
for line in [
    "■ 작성 안내",
    "· 첫 줄은 항목 이름(머리글)이며, 둘째 줄부터 자료를 적습니다.",
    "· 날짜는 2026-01-05 처럼 적거나 엑셀 날짜 서식을 씁니다.",
    "· 금액·수량은 숫자만 적습니다 (쉼표·원 표시는 자동으로 걷어냅니다).",
    "· 구분·분류·기탁처유형처럼 값이 몇 가지로 정해진 열은 자동으로 '범주'로 인식되어 그래프가 됩니다.",
    "· 시트는 필요한 만큼 추가할 수 있으며, 분석 화면에서 시트를 골라 봅니다.",
]:
    g.append([line])
g.column_dimensions["A"].width = 90
p2 = os.path.join(OUT, "기부물품_입력양식.xlsx")
tw.save(p2)
print("saved:", p1, "(rows:", len(rows), ")")
print("saved:", p2)
